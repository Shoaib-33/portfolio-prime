"""Buyers & outreach: imports the Kadaster/BAG ownership exports (owners, owner_properties, properties) and contact lists,
builds a buying profile per company, ranks companies as potential buyers for a listing and renders/sends the emails."""
from __future__ import annotations
import csv, hashlib, hmac, io, math, os, re, smtplib, statistics, threading, unicodedata
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from html import escape


from .. import db
from ..db import T, and_, select
from ..common import _key, bad, log_rule, now, sha

# ---------------- Places ----------------
CITY_ALIAS = {"'s-gravenhage": 'den haag', 's-gravenhage': 'den haag', 'the hague': 'den haag', "'s-hertogenbosch": 'den bosch', 's-hertogenbosch': 'den bosch'}
PROVINCE_OF = {k.lower(): p for p, cities in {
    'Noord-Holland': 'Amsterdam Haarlem Zaandam Hilversum Alkmaar Amstelveen Hoofddorp Haarlemmermeer Purmerend Hoorn Den Helder Velsen IJmuiden Velsen-Noord Bussum Huizen Heemstede Diemen Zandvoort Laren Blaricum Weesp Uithoorn Aalsmeer Badhoevedorp Beverwijk Heerhugowaard Enkhuizen',
    'Zuid-Holland': 'Rotterdam Den_Haag Leiden Dordrecht Zoetermeer Delft Schiedam Vlaardingen Gouda Alphen_aan_den_Rijn Rijswijk Wassenaar Capelle_aan_den_IJssel Spijkenisse Maassluis Leidschendam Voorburg Katwijk Noordwijk Barendrecht Ridderkerk Gorinchem Papendrecht Berkel_en_Rodenrijs Pijnacker Nootdorp Westland Naaldwijk Hellevoetsluis Krimpen_aan_den_IJssel Zwijndrecht Leiderdorp Oegstgeest Lisse',
    'Utrecht': 'Utrecht Amersfoort Veenendaal Nieuwegein Zeist Houten IJsselstein Woerden Soest Leusden Baarn Eemnes Hoogland Hooglanderveen Vleuten De_Meern Maarssen Bunschoten Spakenburg Mijdrecht Breukelen Bilthoven De_Bilt Rhenen Wijk_bij_Duurstede Driebergen Doorn Achterveld',
    'Gelderland': 'Arnhem Nijmegen Apeldoorn Ede Barneveld Nijkerk Harderwijk Doetinchem Zutphen Tiel Wageningen Culemborg Ermelo Voorthuizen Kootwijkerbroek Garderen Stroe Putten Nijkerkerveen Hoevelaken Lunteren Zevenaar Elst Duiven Wijchen Epe Winterswijk Bennekom Renkum Oosterbeek Velp Rheden De_Glind Terschuur',
    'Noord-Brabant': 'Eindhoven Tilburg Breda Den_Bosch Helmond Oss Roosendaal Bergen_op_Zoom Waalwijk Veldhoven Uden Oosterhout Best Boxtel Vught Geldrop Valkenswaard Etten-Leur',
    'Flevoland': 'Almere Lelystad Dronten Zeewolde Emmeloord',
    'Overijssel': 'Enschede Zwolle Deventer Hengelo Almelo Kampen Oldenzaal Raalte',
    'Groningen': 'Groningen Veendam Hoogezand Delfzijl', 'Friesland': 'Leeuwarden Sneek Heerenveen Drachten Harlingen',
    'Drenthe': 'Assen Emmen Hoogeveen Meppel', 'Limburg': 'Maastricht Venlo Heerlen Sittard Roermond Geleen Weert Kerkrade', 'Zeeland': 'Middelburg Vlissingen Goes Terneuzen'}.items()
    for k in [c.replace('_', ' ') for c in cities.split()]}


def city_key(c: str | None) -> str:
    c = re.sub(r'\s+', ' ', (c or '').strip().lower())
    return CITY_ALIAS.get(c, c)


def province_for(city: str | None, learned: dict | None = None) -> str | None:
    k = city_key(city)
    return (learned or {}).get(k) or PROVINCE_OF.get(k)


# ---------------- Reading files ----------------
def _norm(h) -> str:
    s = unicodedata.normalize('NFKD', str(h or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)


def read_table(name: str, data: bytes) -> tuple[list[str], list[list]]:
    """Header row and data rows of the first sheet of an .xlsx file, or of a .csv file."""
    if name.lower().endswith('.csv'):
        text = data.decode('utf-8-sig', errors='replace')
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=',;\t') if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        try:
            from python_calamine import CalamineWorkbook  # fast (Rust); openpyxl is the fallback
            rows = CalamineWorkbook.from_filelike(io.BytesIO(data)).get_sheet_by_index(0).to_python()
        except ImportError:
            from openpyxl import load_workbook
            try: wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            except Exception: raise bad(f'{name}: not a readable Excel file (.xlsx).')
            ws = wb.worksheets[0]; rows = [list(r) for r in ws.iter_rows(values_only=True)]; wb.close()
        except Exception:
            raise bad(f'{name}: not a readable Excel file (.xlsx).')
    rows = [r for r in rows if any(v not in (None, '') for v in r)]
    if not rows: raise bad(f'{name}: the file is empty.')
    return [str(h or '').strip() for h in rows[0]], rows[1:]


CONTACT_COLS = {'company': ('bedrijf', 'bedrijfsnaam', 'company', 'companyname', 'organisatie', 'organisation', 'organization', 'eigenaar', 'koper'),
                'name': ('naam', 'name', 'contactpersoon', 'contact', 'contactname', 'volledigenaam', 'fullname'), 'role': ('functie', 'role', 'rol', 'title', 'jobtitle'),
                'email': ('email', 'emailadres', 'mail', 'emailaddress'), 'phone': ('telefoon', 'telefoonnummer', 'phone', 'tel', 'mobiel', 'mobile'),
                'language': ('taal', 'language', 'lang'), 'notes': ('notities', 'notes', 'opmerkingen')}


def detect(headers: list[str]) -> str | None:
    h = {_norm(x) for x in headers}
    if 'bagpandid' in h and ('factsheet' in h or 'straat' in h or 'gebruiksdoel' in h): return 'properties'
    if 'eigenaar' in h and ('aantalpanden' in h or 'kadasterentiteiten' in h): return 'owners'
    if 'eigenaar' in h and 'bagpandid' in h: return 'holdings'
    if h & set(CONTACT_COLS['email']) and h & set(CONTACT_COLS['company']): return 'contacts'
    return None


def _txt(v) -> str | None:
    if v is None: return None
    if isinstance(v, float) and v.is_integer(): v = int(v)
    s = str(v).strip(); return s or None


def _int(v) -> int | None:
    try: return None if v in (None, '') else int(round(float(str(v).replace(',', '.'))))
    except ValueError: return None


def _float(v) -> float | None:
    try: return None if v in (None, '') else float(str(v).replace(',', '.'))
    except ValueError: return None


def _date(v) -> date | None:
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    s = _txt(v)
    if not s: return None
    for f in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S'):
        try: return datetime.strptime(s[:19], f).date()
        except ValueError: pass
    return None


def bag_id(v) -> str | None:
    s = _txt(v)
    if not s: return None
    s = re.sub(r'\D', '', s.split('.')[0])
    return s.zfill(16) if s else None


def _amount(tok: str) -> int | None:
    m = re.fullmatch(r'([\d.]+)\s*([km]?)', tok.strip().lower())
    if not m: return None
    return int(round(float(m.group(1)) * {'': 1, 'k': 1e3, 'm': 1e6}[m.group(2)]))


def price_band(v) -> tuple[str | None, int | None, int | None]:
    """'600k - 700k' → (label, 600000, 700000); '0' or empty → unknown; a plain amount → (label, n, n)."""
    s = _txt(v)
    if not s or s == '0': return None, None, None
    parts = [p for p in re.split(r'\s*-\s*', s) if p]
    vals = [_amount(p) for p in parts]
    if not vals or None in vals: return s, None, None
    lo, hi = vals[0], vals[-1]
    label = s if len(parts) > 1 else money_short(lo)
    return label, lo, hi


def money_short(n: float | None) -> str:
    if not n: return '—'
    return f'€{n / 1e6:.1f}m'.replace('.0m', 'm') if n >= 1e6 else f'€{round(n / 1e3)}k'


def name_key(s: str) -> str:
    return re.sub(r'\s+', ' ', s.strip().lower())


# ---------------- Imports ----------------
def _companies_by_key() -> dict[str, dict]:
    out = {}
    for c in db.all_(T.buyer_companies):
        out[c['name_key']] = c
        for e in c['legal_entities'] or []: out.setdefault(name_key(e), c)
    return out


def _company(name: str, cache: dict, source: str, group: str | None = 'bedrijf') -> dict:
    k = name_key(name)
    if k not in cache:
        cache[k] = db.ins(T.buyer_companies, {'name': name.strip(), 'name_key': k, 'buyer_group': group, 'source_file': source})
    return cache[k]


def _upsert(table, rows: list[dict], keys: list[str]) -> tuple[int, int]:
    """Bulk upsert: COPY into a temporary table, then INSERT … ON CONFLICT DO UPDATE. Returns (created, updated)."""
    if not rows: return 0, 0
    cols = list(rows[0]); q = lambda c: '"%s"' % c
    sets = ', '.join(f'{q(c)} = EXCLUDED.{q(c)}' for c in cols if c not in keys)
    raw = db.engine().raw_connection()
    try:
        conn = raw.driver_connection
        with conn.transaction(), conn.cursor() as cur:
            cur.execute(f'CREATE TEMP TABLE _up (LIKE "{table.name}" INCLUDING DEFAULTS) ON COMMIT DROP')
            with cur.copy(f'COPY _up ({", ".join(map(q, cols))}) FROM STDIN') as cp:
                for r in rows: cp.write_row([r[c] for c in cols])
            cur.execute(f'INSERT INTO "{table.name}" ({", ".join(map(q, cols))}) SELECT {", ".join(map(q, cols))} FROM _up ON CONFLICT ({", ".join(map(q, keys))}) DO UPDATE SET {sets} RETURNING (xmax = 0)')
            flags = [r[0] for r in cur.fetchall()]
    finally:
        raw.close()
    return sum(flags), len(flags) - sum(flags)


def _row_get(headers: list[str]):
    idx = {}
    for i, h in enumerate(headers): idx.setdefault(_norm(h), i)
    def get(row, *names, prefix=False):
        for n in names:
            if prefix:
                i = next((j for k, j in idx.items() if k.startswith(n)), None)
            else:
                i = idx.get(n)
            if i is not None and i < len(row): return row[i]
        return None
    return get


def import_properties(headers, rows, fname) -> dict:
    g = _row_get(headers); out = {}; skipped = 0
    for r in rows:
        bid = bag_id(g(r, 'bagpandid'))
        if not bid: skipped += 1; continue
        label, lo, hi = price_band(g(r, 'koopsombandbreedte', 'koopsom'))
        street, num = _txt(g(r, 'straat')), _txt(g(r, 'nummer'))
        number = ''.join(x for x in [num, _txt(g(r, 'letter'))] if x) + (f"-{_txt(g(r, 'toevoeging'))}" if _txt(g(r, 'toevoeging')) else '') if num else None
        owner = _txt(g(r, 'eigenaar'))
        out[bid] = {'bag_pand_id': bid, 'street': street, 'number': number, 'postcode': _txt(g(r, 'postcode')), 'city': _txt(g(r, 'plaats')), 'municipality': _txt(g(r, 'gemeente')), 'province': _txt(g(r, 'provincie')),
                    'multiple_addresses': bool(_txt(g(r, 'meerdereadressen'))), 'use_purpose': _txt(g(r, 'gebruiksdoel')), 'floor_m2': _int(g(r, 'gebruiksoppervlaktem2', 'gebruiksoppervlakte')),
                    'plot_m2': _int(g(r, 'perceeloppervlaktem2', 'perceeloppervlakte')), 'build_year': _int(g(r, 'bouwjaar')), 'monument': (_txt(g(r, 'monument')) or '').lower() in ('ja', 'yes', 'true', '1'),
                    'status': _txt(g(r, 'status')), 'energy_labels': _txt(g(r, 'energielabels', 'energielabel')), 'occupants': _int(g(r, 'aantaleindgebruikers')), 'purchase_date': _date(g(r, 'aankoopdatum')),
                    'years_owned': _float(g(r, 'jareninbezit')), 'segment': _txt(g(r, 'segment')), 'price_band': label, 'price_low': lo, 'price_high': hi, 'transaction_type': _txt(g(r, 'typekooptransactie')),
                    'buyer_group': _txt(g(r, 'kopergroep')), 'owner_name': None if not owner or 'natuurlijk persoon' in owner.lower() else owner, 'transaction_address': _txt(g(r, 'transactieadres')),
                    'factsheet': _txt(g(r, 'factsheet')), 'source_file': _txt(g(r, 'bronbestand')) or fname, 'updated_at': now()}
    created, updated = _upsert(T.buyer_buildings, list(out.values()), ['bag_pand_id'])
    return {'rows': len(rows), 'created': created, 'updated': updated, 'skipped': skipped + len(rows) - skipped - len(out)}


def import_owners(headers, rows, fname) -> dict:
    g = _row_get(headers); cache = _companies_by_key(); created = updated = skipped = 0
    for r in rows:
        name = _txt(g(r, 'eigenaar'))
        if not name: skipped += 1; continue
        ents = [e.strip() for e in re.split(r'[;|\n]', _txt(g(r, 'kadasterentiteiten')) or '') if e.strip()]
        if len(ents) == 1 and ', ' in ents[0] and ents[0].count(',') > 1: ents = [e.strip() for e in ents[0].split(',') if e.strip()]
        vals = {'name': name, 'buyer_group': _txt(g(r, 'kopergroep')) or 'bedrijf', 'legal_entities': ents, 'buildings': _int(g(r, 'aantalpanden')) or 0, 'total_m2': _int(g(r, 'totaalm2')) or 0,
                'first_purchase': _date(g(r, 'eersteaankoop')), 'last_purchase': _date(g(r, 'laatsteaankoop')), 'avg_years': _float(g(r, 'gemjareninbezit')),
                'recent_count': _int(g(r, 'aantalrecent', prefix=True)) or 0, 'long_count': _int(g(r, 'aantallang', prefix=True)) or 0, 'segment': _txt(g(r, 'segment')),
                'cities': [c.strip() for c in (_txt(g(r, 'plaatsen')) or '').split(',') if c.strip()], 'source_file': fname, 'updated_at': now()}
        c = cache.get(name_key(name))
        if c: db.upd(T.buyer_companies, T.buyer_companies.c.id == c['id'], vals); updated += 1
        else: cache[name_key(name)] = db.ins(T.buyer_companies, {**vals, 'name_key': name_key(name)}); created += 1
    return {'rows': len(rows), 'created': created, 'updated': updated, 'skipped': skipped}


def import_holdings(headers, rows, fname) -> dict:
    g = _row_get(headers); cache = _companies_by_key(); out = {}; skipped = 0; before = db.count(T.buyer_companies)
    for r in rows:
        name, bid = _txt(g(r, 'eigenaar')), bag_id(g(r, 'bagpandid'))
        if not name or not bid or 'natuurlijk persoon' in name.lower(): skipped += 1; continue
        c = _company(name, cache, fname)
        out[(c['id'], bid)] = {'company_id': c['id'], 'bag_pand_id': bid, 'address': _txt(g(r, 'adres')), 'city': _txt(g(r, 'plaats')), 'purchase_date': _date(g(r, 'aankoopdatum')),
                               'years_owned': _float(g(r, 'jareninbezit')), 'segment': _txt(g(r, 'segment')), 'source_file': _txt(g(r, 'bronbestand')) or fname}
    created, updated = _upsert(T.buyer_holdings, list(out.values()), ['company_id', 'bag_pand_id'])
    return {'rows': len(rows), 'created': created, 'updated': updated, 'skipped': skipped, 'newCompanies': db.count(T.buyer_companies) - before}


EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')


def import_contacts(headers, rows, fname, by: str | None) -> dict:
    g = _row_get(headers); cache = _companies_by_key(); created = updated = skipped = 0; problems = []; new_companies = 0
    existing = {(c['company_id'], (c['email'] or '').lower()): c for c in db.all_(T.buyer_contacts)}
    for n, r in enumerate(rows, start=2):
        company, email = _txt(g(r, *CONTACT_COLS['company'])), (_txt(g(r, *CONTACT_COLS['email'])) or '').lower()
        name = _txt(g(r, *CONTACT_COLS['name'])) or (email.split('@')[0] if email else None)
        if not company or not email: skipped += 1; problems.append(f'Row {n}: company and email are required'); continue
        if not EMAIL_RE.match(email): skipped += 1; problems.append(f'Row {n}: “{email}” is not a valid email'); continue
        if name_key(company) not in cache: new_companies += 1
        c = _company(company, cache, fname, None)
        lang = (_txt(g(r, *CONTACT_COLS['language'])) or 'nl').lower()[:2]
        vals = {'name': name[:120], 'role': _txt(g(r, *CONTACT_COLS['role'])), 'email': email, 'phone': _txt(g(r, *CONTACT_COLS['phone'])), 'language': 'en' if lang in ('en', 'de', 'fr') else 'nl', 'notes': _txt(g(r, *CONTACT_COLS['notes']))}
        old = existing.get((c['id'], email))
        if old: db.upd(T.buyer_contacts, T.buyer_contacts.c.id == old['id'], vals); updated += 1
        else: existing[(c['id'], email)] = db.ins(T.buyer_contacts, {**vals, 'company_id': c['id'], 'created_by': by}); created += 1
    return {'rows': len(rows), 'created': created, 'updated': updated, 'skipped': skipped, 'newCompanies': new_companies, 'problems': problems[:20]}


def import_file(fname: str, data: bytes | None, by: str | None, table: tuple | None = None) -> dict:
    headers, rows = table or read_table(fname, data); kind = detect(headers)
    if not kind:
        raise bad(f'{fname}: columns not recognised ({", ".join(h for h in headers[:8] if h)}…). Expected the owners, owner_properties or properties export, or a contacts list with company and email columns.')
    rep = {'properties': import_properties, 'owners': import_owners, 'holdings': import_holdings}[kind](headers, rows, fname) if kind != 'contacts' else import_contacts(headers, rows, fname, by)
    return {'file': fname, 'kind': kind, **rep}


ORDER = {'properties': 0, 'owners': 1, 'holdings': 2, 'contacts': 3}
_import_lock = threading.Lock()


def run_batch(batch: list[tuple[str, str, bytes]], me: dict):
    """Background import of one upload: (import row id, file name, bytes) per file. Buildings first, then owners and holdings,
    contacts last, so one upload of all files gives complete profiles. Each file's row in buyer_imports shows progress."""
    I = T.buyer_imports
    with _import_lock:
        parsed = []
        for iid, name, data in batch:
            try:
                headers, rows = read_table(name, data); kind = detect(headers)
                if not kind: raise bad(f'Columns not recognised ({", ".join(h for h in headers[:6] if h)}…). Expected the owners, owner_properties or properties export, or a contacts list with company and email columns.')
                db.upd(I, I.c.id == iid, {'kind': kind, 'rows': len(rows)}); parsed.append((ORDER[kind], iid, name, (headers, rows)))
            except Exception as e:
                db.upd(I, I.c.id == iid, {'status': 'failed', 'error': getattr(e, 'message', None) or str(e)[:300]})
        done = []
        for _, iid, name, table in sorted(parsed, key=lambda x: x[0]):
            db.upd(I, I.c.id == iid, {'status': 'processing'})
            try:
                r = import_file(name, None, me['id'], table); done.append(r)
                db.upd(I, I.c.id == iid, {'status': 'done', 'created': r['created'], 'updated': r['updated'], 'skipped': r['skipped'], 'report': r})
            except Exception as e:
                db.upd(I, I.c.id == iid, {'status': 'failed', 'error': getattr(e, 'message', None) or str(e)[:300]})
        if any(r['kind'] != 'contacts' for r in done): refresh_profiles()
        if done:
            from ..common import audit
            audit(None, 'buyers_import', 'buyer_import', None, {'files': ', '.join(r['file'] for r in done), 'rows': sum(r['rows'] for r in done)}, actor=me)
            log_rule('buyers', f"{me['name']} imported " + ', '.join('%s (%s, %s rows)' % (r['file'], r['kind'], format(r['rows'], ',')) for r in done))


def start_batch(batch, me):
    threading.Thread(target=run_batch, args=(batch, me), name='pp-import', daemon=True).start()


# ---------------- Buying profile per company ----------------
USE_GROUPS = [('residential', 'woonfunctie'), ('retail', 'winkelfunctie'), ('office', 'kantoorfunctie'), ('industrial', 'industriefunctie'), ('hospitality', 'logiesfunctie'),
              ('healthcare', 'gezondheidszorgfunctie'), ('education', 'onderwijsfunctie'), ('assembly', 'bijeenkomstfunctie')]


def _use_groups(purpose: str | None) -> list[str]:
    if not purpose: return []
    p = purpose.lower(); gs = [k for k, w in USE_GROUPS if w in p]
    return gs or ['other']


def refresh_profiles() -> int:
    """Recomputes every company's profile from its holdings joined with the building data."""
    B, H = T.buyer_buildings, T.buyer_holdings
    learned = {}
    for r in db.rows(select(B.c.city, B.c.province).where(B.c.city.isnot(None), B.c.province.isnot(None)).distinct()): learned.setdefault(city_key(r['city']), r['province'])
    rows = db.rows(select(H.c.company_id, H.c.bag_pand_id, H.c.city.label('h_city'), H.c.purchase_date.label('h_date'), H.c.segment.label('h_seg'),
                          B.c.city, B.c.province, B.c.use_purpose, B.c.floor_m2, B.c.price_low, B.c.price_high, B.c.price_band, B.c.transaction_type, B.c.purchase_date, B.c.build_year)
                   .select_from(H.outerjoin(B, H.c.bag_pand_id == B.c.bag_pand_id)))
    by: dict[str, list] = {}
    for r in rows: by.setdefault(r['company_id'], []).append(r)
    today = date.today(); five = today - timedelta(days=round(5 * 365.25)); one = today - timedelta(days=365)
    companies = db.all_(T.buyer_companies)
    for c in companies:
        hs = by.get(c['id'], [])
        cities, provinces, recent_cities, uses = {}, {}, {}, {}
        floors, years, deals, invest, develop, recent5, recent1, known = [], [], {}, 0, 0, 0, 0, 0
        dates = []
        for h in hs:
            city = h['h_city'] or h['city']; ck = city_key(city)
            d = h['h_date'] or h['purchase_date']
            if ck:
                cities[ck] = cities.get(ck, 0) + 1
                pv = h['province'] or province_for(city, learned)
                if pv: provinces[pv] = provinces.get(pv, 0) + 1
            if d:
                dates.append(d)
                if d >= five:
                    recent5 += 1
                    if ck: recent_cities[ck] = recent_cities.get(ck, 0) + 1
                if d >= one: recent1 += 1
            gs = _use_groups(h['use_purpose'])
            if gs: known += 1
            for gk in gs: uses[gk] = uses.get(gk, 0) + 1
            if h['floor_m2']: floors.append(h['floor_m2'])
            if h['build_year']: years.append(h['build_year'])
            if h['price_high']: deals[(d, h['price_band'])] = ((h['price_low'] or 0) + h['price_high']) / 2  # one transaction can cover many buildings
            t = (h['transaction_type'] or '').lower()
            if 'belegging' in t: invest += 1
            if 'ontwikkeling' in t or 'forward funding' in t or 'transformatie' in t: develop += 1
        sizes = sorted(deals.values())
        names = {}  # display names for city keys
        for h in hs:
            city = h['h_city'] or h['city']
            if city: names.setdefault(city_key(city), city)
        profile = {'holdings': len(hs), 'cityCounts': cities, 'cityNames': names, 'provinceCounts': provinces, 'recentCityCounts': recent_cities, 'uses': uses, 'usesKnown': known,
                   'residentialShare': round(uses.get('residential', 0) / known, 3) if known else None, 'medianFloorM2': round(statistics.median(floors)) if floors else None,
                   'medianBuildYear': round(statistics.median(years)) if years else None, 'deals': len(sizes), 'medianDeal': round(statistics.median(sizes)) if sizes else None,
                   'maxDeal': round(sizes[-1]) if sizes else None, 'investmentDeals': invest, 'developmentDeals': develop, 'recent5y': recent5, 'recent12m': recent1}
        vals = {'profile': profile}
        if not c['source_file'] or 'owner' not in (c['source_file'] or '').lower() or not c['buildings']:  # no owners-file summary: derive it
            vals.update({'buildings': max(c['buildings'] or 0, len(hs)), 'first_purchase': min(dates) if dates else c['first_purchase'], 'last_purchase': max(dates) if dates else c['last_purchase'],
                         'recent_count': recent5, 'cities': [names[k] for k, _ in sorted(cities.items(), key=lambda x: -x[1])][:30] or c['cities']})
        db.upd(T.buyer_companies, T.buyer_companies.c.id == c['id'], vals)
    return len(companies)


# ---------------- Matching: who could buy this listing? ----------------
def _plural(n, one, many): return f'{n:,} {one if n == 1 else many}'


def match(listing: dict, prop: dict, companies: list[dict], lang: str = 'en') -> list[dict]:
    nl = lang == 'nl'
    city = prop['city']; ck = city_key(city); prov = province_for(city)
    units = max(1, listing.get('_units') or 1); price = listing['asking_price'] or 0
    out = []
    for c in companies:
        p = c['profile'] or {}; reasons = []; score = 0.0
        n_city = (p.get('cityCounts') or {}).get(ck, 0); r_city = (p.get('recentCityCounts') or {}).get(ck, 0); n_prov = (p.get('provinceCounts') or {}).get(prov, 0) if prov else 0
        if not n_city and city and any(city_key(x) == ck for x in c['cities'] or []): n_city = 1
        if n_city:
            score += min(35, 20 + 5 * math.log2(1 + n_city))
            reasons.append((f'Bezit {_plural(n_city, "pand", "panden")} in {city}' if nl else f'Owns {_plural(n_city, "building", "buildings")} in {city}') + (f' ({r_city} {"recent gekocht" if nl else "bought in the last 5 years"})' if r_city else ''))
        elif n_prov:
            score += min(18, 10 + 2 * math.log2(1 + n_prov))
            reasons.append(f'Bezit {_plural(n_prov, "pand", "panden")} in de provincie {prov}' if nl else f'Owns {_plural(n_prov, "building", "buildings")} in {prov} province')
        share = p.get('residentialShare')
        if share is not None:
            score += 20 * share
            if share >= 0.5: reasons.append(f'{round(share * 100)}% woningen in portefeuille' if nl else f'{round(share * 100)}% of the portfolio is residential')
        else: score += 6
        mx, md = p.get('maxDeal'), p.get('medianDeal')
        if mx and price:
            if mx >= price * 0.7: score += 20; reasons.append(f'Deed transacties tot {money_short(mx)}' if nl else f'Has done deals up to {money_short(mx)}')
            elif mx >= price * 0.4: score += 12; reasons.append(f'Grootste transactie {money_short(mx)}' if nl else f'Largest deal {money_short(mx)}')
            else: score += 4
        else: score += 6
        b = c['buildings'] or 0
        score += 10 if b >= units else 6 if b * 3 >= units else 2
        if units > 1 and b >= units: reasons.append(f'Beheert {b:,} panden: kan een portefeuille van {units} eenheden dragen' if nl else f'Holds {b:,} buildings: can absorb a {units}-unit portfolio')
        last = c['last_purchase']
        if last:
            age = (date.today() - last).days
            score += 10 if age <= 365 else 7 if age <= 3 * 365 else 4 if age <= 5 * 365 else 0
            if age <= 3 * 365: reasons.append(f'Laatste aankoop {last.strftime("%m-%Y")}' if nl else f'Last purchase {last.strftime("%b %Y")}')
        if p.get('investmentDeals'): score += 3
        if (p.get('recent5y') or c['recent_count'] or 0) >= 10: score += 2
        out.append({'id': c['id'], 'name': c['name'], 'score': min(100, round(score)), 'reasons': reasons[:4], 'buildings': b, 'lastPurchase': last, 'segment': c['segment'], 'stage': c['stage'],
                    'doNotContact': c['do_not_contact'], 'inCity': n_city})
    return sorted(out, key=lambda x: (-x['score'], -x['buildings']))


# ---------------- Emails ----------------
TEMPLATES = {
    'nl': {'subject': '{property} in {city}: besloten aanbod voor {company}',
           'body': 'Beste {first_name},\n\nPortfolio & Prime brengt een nieuwe investering in de markt die aansluit bij de portefeuille van {company}: {property} in {city} ({type}).\n\nWaarom wij aan u dachten: {reason}.\n\nHet aanbod is besloten. Via uw persoonlijke link bekijkt u de details, de cijfers en de stukken:\n{link}\n\nIk licht het graag persoonlijk toe.\n\nMet vriendelijke groet,\n{sender}\nPortfolio & Prime'},
    'en': {'subject': '{property} in {city}: a private opportunity for {company}',
           'body': 'Dear {first_name},\n\nPortfolio & Prime is bringing a new investment to market that fits the portfolio of {company}: {property} in {city} ({type}).\n\nWhy we thought of you: {reason}.\n\nThe offering is private. Your personal link opens the details, the numbers and the documents:\n{link}\n\nI would be glad to walk you through it.\n\nKind regards,\n{sender}\nPortfolio & Prime'}}
FOOTER = {'nl': 'U ontvangt deze e-mail omdat {company} actief is op de Nederlandse vastgoedmarkt. Geen berichten meer ontvangen? {unsubscribe}',
          'en': 'You receive this email because {company} is active in the Dutch property market. No more emails? {unsubscribe}'}


def personalise(text_: str, values: dict) -> str:
    """Fills {placeholders}; {link} and {unsubscribe} stay for the send step (they carry the personal token)."""
    return re.sub(r'\{(\w+)\}', lambda m: str(values.get(m.group(1), m.group(0))) if m.group(1) not in ('link', 'unsubscribe') else m.group(0), text_)


def values_for(listing: dict, prop: dict, units: int, company: dict, contact: dict | None, reason: str, sender: str, lang: str) -> dict:
    first = (contact['name'].split()[0] if contact and contact['name'] else ('heer/mevrouw' if lang == 'nl' else 'Sir or Madam'))
    typ = prop['property_type_nl'] if lang == 'nl' and prop.get('property_type_nl') else prop['property_type']
    if units > 1: typ += f', {units} {"eenheden" if lang == "nl" else "units"}'
    return {'first_name': first, 'name': contact['name'] if contact else first, 'company': company['name'], 'property': (listing.get('name_nl') if lang == 'nl' and listing.get('name_nl') else listing['name']),
            'city': prop['city'], 'type': typ, 'reason': reason[:1].lower() + reason[1:] if reason else ('uw activiteit op de Nederlandse woningmarkt' if lang == 'nl' else 'your activity in the Dutch residential market'),
            'sender': sender, 'price': 'op aanvraag' if lang == 'nl' and listing['price_on_request'] else 'on request' if listing['price_on_request'] else f"€{listing['asking_price']:,}".replace(',', '.')}


def token_for(message_id: str) -> str:
    """Personal token in the email links, derived from the message id (nothing secret is stored)."""
    return hmac.new(_key(), f'outreach:{message_id}'.encode(), hashlib.sha256).hexdigest()[:40]


def message_by_token(token: str) -> dict | None:
    return db.get(T.outreach_messages, token_hash=sha(token)) if re.fullmatch(r'[0-9a-f]{40}', token or '') else None


def links(base: str, token: str) -> dict:
    return {'link': f'{base}/api/m/{token}', 'unsubscribe': f'{base}/u/{token}', 'pixel': f'{base}/api/m/{token}/o.gif'}


def render(m: dict, campaign: dict, lang: str) -> tuple[str, str, str, dict]:
    """Final subject, plain text and HTML of one message."""
    t = token_for(m['id']); L = links(campaign['base_url'], t)
    company = db.get(T.buyer_companies, id=m['company_id'])
    body = m['body'].replace('{link}', L['link']).replace('{unsubscribe}', L['unsubscribe'])
    foot = FOOTER[lang].replace('{company}', company['name'] if company else '').replace('{unsubscribe}', L['unsubscribe'])
    text_ = f'{body}\n\n--\n{foot}'
    paras = ''.join(f'<p style="margin:0 0 14px">{escape(p).replace(chr(10), "<br>")}</p>' for p in body.split('\n\n'))
    paras = paras.replace(escape(L['link']), f'<a href="{escape(L["link"])}" style="color:#1f3a5f;font-weight:600">{"Bekijk het aanbod" if lang == "nl" else "View the opportunity"}</a>')
    foot_html = escape(FOOTER[lang].replace('{company}', company['name'] if company else '')).replace('{unsubscribe}', '<a href="%s" style="color:#6b7280">%s</a>' % (escape(L['unsubscribe']), 'Afmelden' if lang == 'nl' else 'Unsubscribe'))
    html = ('<div style="font-family:Georgia,serif;font-size:15px;line-height:1.55;color:#1d232b;max-width:560px">' + paras +
            '<p style="margin:26px 0 0;font:12px/1.5 Arial,sans-serif;color:#6b7280">' + foot_html + '</p>'
            '<img src="%s" width="1" height="1" alt=""></div>' % escape(L['pixel']))
    return m['subject'], text_, html, L


def mail_mode() -> dict:
    host = os.environ.get('SMTP_HOST')
    return {'mode': 'smtp' if host else 'simulated', 'from': os.environ.get('MAIL_FROM') or 'Portfolio & Prime <no-reply@portfolio-prime.nl>', 'host': host,
            'perTick': int(os.environ.get('PP_MAIL_BATCH') or 25)}


def _smtp_send(to: str, subject: str, text_: str, html: str, reply_to: str | None, unsub: str):
    cfg = mail_mode(); msg = EmailMessage()
    msg['From'] = cfg['from']; msg['To'] = to; msg['Subject'] = subject
    if reply_to: msg['Reply-To'] = reply_to
    msg['List-Unsubscribe'] = f'<{unsub}>'; msg['List-Unsubscribe-Post'] = 'List-Unsubscribe=One-Click'
    msg.set_content(text_); msg.add_alternative(html, subtype='html')
    port = int(os.environ.get('SMTP_PORT') or 587)
    with (smtplib.SMTP_SSL(cfg['host'], port, timeout=30) if port == 465 else smtplib.SMTP(cfg['host'], port, timeout=30)) as s:
        if port != 465: s.starttls()
        if os.environ.get('SMTP_USER'): s.login(os.environ['SMTP_USER'], os.environ.get('SMTP_PASSWORD', ''))
        s.send_message(msg)


def send_queued(limit: int | None = None) -> dict:
    """Sends up to `limit` queued messages (called by the background job every 30 s and by “Send now”)."""
    cfg = mail_mode(); M = T.outreach_messages
    queue = db.all_(M, M.c.status == 'queued', order=M.c.created_at, limit=limit or cfg['perTick'])
    sent = failed = 0; camps = {}
    for m in queue:
        camp = camps.get(m['campaign_id']) or db.get(T.outreach_campaigns, id=m['campaign_id']); camps[m['campaign_id']] = camp
        contact = db.get(T.buyer_contacts, id=m['contact_id'])
        if not contact or contact['unsubscribed_at']:
            db.upd(M, M.c.id == m['id'], {'status': 'cancelled', 'error': 'Contact unsubscribed'}); continue
        try:
            subject, text_, html, L = render(m, camp, camp['language'])
            if cfg['mode'] == 'smtp':
                creator = db.get(T.users, id=camp['created_by']) if camp['created_by'] else None
                _smtp_send(m['to_email'], subject, text_, html, creator['email'] if creator else None, L['unsubscribe'])
            else:
                db.ins(T.outbox, {'to_email': m['to_email'], 'subject': subject, 'body': text_, 'kind': 'outreach', 'status': 'simulated'}, returning=False)
            db.upd(M, M.c.id == m['id'], {'status': 'sent', 'sent_at': now(), 'error': None}); sent += 1
            db.upd(T.buyer_companies, and_(T.buyer_companies.c.id == m['company_id'], T.buyer_companies.c.stage == 'new'), {'stage': 'contacted'})
        except Exception as e:
            db.upd(M, M.c.id == m['id'], {'status': 'failed', 'error': str(e)[:300]}); failed += 1
    if sent or failed: log_rule('outreach', f"Outreach: {sent} email{'s' if sent != 1 else ''} {'sent' if cfg['mode'] == 'smtp' else 'sent (simulated)'}{f', {failed} failed' if failed else ''}")
    left = db.count(M, M.c.status == 'queued')
    return {'sent': sent, 'failed': failed, 'queued': left, 'mode': cfg['mode']}


STATUS_ORDER = ['queued', 'sent', 'opened', 'clicked', 'registered']


def advance(m: dict, status: str, field: str):
    """Moves a message forward in the funnel (never backwards) and stamps the time once."""
    vals = {}
    if not m.get(field): vals[field] = now()
    if m['status'] in STATUS_ORDER and STATUS_ORDER.index(status) > STATUS_ORDER.index(m['status']): vals['status'] = status
    if field == 'clicked_at' and not m.get('opened_at'): vals['opened_at'] = now()
    if vals: db.upd(T.outreach_messages, T.outreach_messages.c.id == m['id'], vals)
